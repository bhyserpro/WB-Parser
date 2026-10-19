"""Парсер цен на iPhone 16 с Wildberries.

Выдача берётся из JSON API поиска, которым пользуется сам сайт
(``/__internal/search/exactmatch/...``). API закрыт антиботом WB, поэтому
cookie ``x_wbaas_token`` сначала получается в настоящем браузере
(см. ``wb_token.py``), а дальше страницы запрашиваются обычным ``requests``.
"""

from __future__ import annotations

import argparse
import csv
import json
import logging
import random
import re
import statistics
import time
from collections import Counter
from dataclasses import asdict, dataclass, fields
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

import requests

from wb_token import TokenRejected, TokenStore

log = logging.getLogger("wb_parser")

SITE = "https://www.wildberries.ru"
SEARCH_API = f"{SITE}/__internal/search/exactmatch/ru/common/v18/search"
DEFAULT_DEST = "-1257786"  # регион доставки (Москва); влияет на цены и наличие
RETRY_STATUSES = {429, 500, 502, 503, 504}
TOKEN_STATUSES = {401, 403, 498}

MODELS = ("iPhone 16", "iPhone 16 Plus", "iPhone 16 Pro", "iPhone 16 Pro Max")
MODEL_RE = re.compile(
    r"(?:iphone|айфон)\s*16\s*(pro\s*max|promax|pro|plus|e)?(?![a-zа-я0-9])",
    re.IGNORECASE,
)
MEMORY_RE = re.compile(r"(?<!\d)(\d{3,4}|1|2)\s*(gb|гб|tb|тб)\b", re.IGNORECASE)
VALID_MEMORY = {"128 ГБ", "256 ГБ", "512 ГБ", "1 ТБ"}
ACCESSORY_WORDS = (
    "чехол", "чехл", "стекл", "пленк", "плёнк", "кабел", "провод", "зарядн",
    "зарядк", "адаптер", "блок питания", "накладк", "бампер", "защитн",
    "держател", "подставк", "ремеш", "корпус", "дисплей", "аккумулятор",
    "батаре", "шлейф", "коробк", "муляж", "макет", "наклейк", "линз", "кольц",
    "кошел", "сумк", "брелок", "стилус", "наушник", "airpods", "power bank",
    "повербанк", "павербанк", "переходник", "футляр", "протектор", "запчаст",
    "case", "glass", "cable", "cover", "magsafe",
)
KNOWN_COLORS = (
    "black titanium", "white titanium", "natural titanium", "desert titanium",
    "черный титан", "чёрный титан", "белый титан", "натуральный титан",
    "пустынный титан", "ultramarine", "ультрамарин", "teal", "бирюзов",
    "pink", "розов", "white", "бел", "black", "черн", "чёрн",
)


@dataclass
class Item:
    """Один товар из выдачи."""

    id: int
    name: str
    brand: str
    model: str | None
    memory: str | None
    color: str | None
    price: float | None
    price_basic: float | None
    discount_pct: float | None
    seller: str
    seller_id: int | None
    rating: float | None
    feedbacks: int | None
    url: str
    collected_at: str


# ---------------------------------------------------------------- загрузка


def build_session(tokens: TokenStore) -> requests.Session:
    """Создаёт сессию с cookie и User-Agent того браузера, что прошёл проверку."""
    session = requests.Session()
    tokens.apply(session)
    session.headers.update({
        "Accept": "*/*",
        "Accept-Language": "ru-RU,ru;q=0.9,en;q=0.8",
        "Referer": f"{SITE}/catalog/0/search.aspx",
    })
    return session


def fetch_page(
    session: requests.Session,
    query: str,
    page: int,
    dest: str = DEFAULT_DEST,
    retries: int = 5,
    timeout: float = 20.0,
) -> dict[str, Any]:
    """Загружает одну страницу выдачи и возвращает JSON.

    Повторяет запрос с экспоненциальной паузой при 429/5xx, таймаутах и
    обрывах. При 401/403/498 бросает ``TokenRejected``: токен надо обновить.
    """
    params = {
        "ab_testing": "false", "appType": "1", "curr": "rub", "dest": dest,
        "lang": "ru", "locale": "ru", "page": str(page), "query": query,
        "resultset": "catalog", "sort": "popular", "spp": "30",
        "suppressSpellcheck": "false",
    }
    for attempt in range(retries + 1):
        try:
            resp = session.get(SEARCH_API, params=params, timeout=timeout)
        except (requests.Timeout, requests.ConnectionError) as exc:
            reason, wait = f"{type(exc).__name__}", None
        else:
            if resp.status_code in TOKEN_STATUSES:
                raise TokenRejected(f"HTTP {resp.status_code} на странице {page}")
            if resp.status_code == 200:
                try:
                    return resp.json()
                except ValueError:
                    reason, wait = "ответ не JSON", None
            elif resp.status_code in RETRY_STATUSES:
                reason = f"HTTP {resp.status_code}"
                retry_after = resp.headers.get("Retry-After", "")
                wait = float(retry_after) if retry_after.isdigit() else None
            else:
                resp.raise_for_status()
                reason, wait = f"HTTP {resp.status_code}", None
        if attempt == retries:
            raise RuntimeError(f"Страница {page}: {reason}, попытки кончились")
        wait = wait if wait is not None else 2 ** (attempt + 1) + random.uniform(0, 1)
        log.warning("Страница %d: %s, повтор через %.1f с", page, reason, wait)
        time.sleep(wait)
    raise AssertionError("unreachable")


# ------------------------------------------------------------------ разбор


def _rub(kopecks: Any) -> float | None:
    """Переводит копейки из API в рубли."""
    return round(kopecks / 100, 2) if isinstance(kopecks, (int, float)) and kopecks > 0 else None


def _extract_prices(product: dict[str, Any]) -> tuple[float | None, float | None]:
    """Возвращает (текущая цена, цена без скидки) по самому дешёвому размеру."""
    best: tuple[float, float | None] | None = None
    for size in product.get("sizes") or []:
        price = size.get("price") or {}
        current = _rub(price.get("product")) or _rub(price.get("total"))
        if current is not None and (best is None or current < best[0]):
            best = (current, _rub(price.get("basic")))
    if best:
        return best
    # старый формат выдачи: цены прямо в карточке
    return _rub(product.get("salePriceU")), _rub(product.get("priceU"))


def detect_model(name: str) -> str | None:
    """Определяет модель линейки iPhone 16 по названию."""
    match = MODEL_RE.search(name.replace(" ", " "))
    if not match:
        return None
    suffix = (match.group(1) or "").lower().replace(" ", "")
    return {
        "": "iPhone 16", "plus": "iPhone 16 Plus", "pro": "iPhone 16 Pro",
        "promax": "iPhone 16 Pro Max", "e": "iPhone 16e",
    }[suffix]


def detect_memory(*texts: str) -> str | None:
    """Ищет объём памяти (128/256/512 ГБ, 1 ТБ) в переданных строках."""
    for text in texts:
        for number, unit in MEMORY_RE.findall(text or ""):
            value = f"{int(number)} {'ТБ' if unit.lower() in ('tb', 'тб') else 'ГБ'}"
            if value in VALID_MEMORY:
                return value
    return None


def detect_color(product: dict[str, Any]) -> str | None:
    """Берёт цвет из поля colors, иначе ищет известный цвет в названии."""
    colors = [c.get("name") for c in product.get("colors") or [] if c.get("name")]
    if colors:
        return ", ".join(colors)
    name = (product.get("name") or "").lower()
    for color in KNOWN_COLORS:
        if color in name:
            return color
    return None


def parse_items(data: dict[str, Any], collected_at: str) -> list[Item]:
    """Превращает JSON страницы выдачи в список ``Item``."""
    products = data.get("products")
    if products is None:
        products = (data.get("data") or {}).get("products") or []
    items: list[Item] = []
    for p in products:
        if "id" not in p:
            continue
        name = (p.get("name") or "").strip()
        price, basic = _extract_prices(p)
        size_names = " ".join(
            f"{s.get('name', '')} {s.get('origName', '')}" for s in p.get("sizes") or []
        )
        discount = round((1 - price / basic) * 100, 1) if price and basic and basic > price else None
        items.append(Item(
            id=int(p["id"]),
            name=name,
            brand=p.get("brand") or "",
            model=detect_model(name),
            memory=detect_memory(name, size_names),
            color=detect_color(p),
            price=price,
            price_basic=basic,
            discount_pct=discount,
            seller=p.get("supplier") or "",
            seller_id=p.get("supplierId"),
            rating=p.get("reviewRating") or p.get("rating"),
            feedbacks=p.get("nmFeedbacks") if p.get("nmFeedbacks") is not None else p.get("feedbacks"),
            url=f"{SITE}/catalog/{p['id']}/detail.aspx",
            collected_at=collected_at,
        ))
    return items


# ---------------------------------------------------------- фильтр и дедуп


def is_accessory(item: Item) -> bool:
    """True, если по названию это аксессуар, а не смартфон."""
    name = item.name.lower()
    return any(word in name for word in ACCESSORY_WORDS)


def filter_items(
    items: Iterable[Item], include_accessories: bool = False, min_price: float = 20000
) -> list[Item]:
    """Оставляет только смартфоны iPhone 16 / Plus / Pro / Pro Max в наличии.

    Признаки смартфона: модель из линейки в названии, нет слов-аксессуаров,
    цена не ниже ``min_price`` (аксессуары без стоп-слов отсекаются ценой).
    """
    items = list(items)
    if include_accessories:
        return items
    kept = [
        it for it in items
        if it.model in MODELS
        and not is_accessory(it)
        and it.price is not None
        and it.price >= min_price
    ]
    log.info("Фильтр: оставлено %d из %d", len(kept), len(items))
    return kept


def dedupe_items(items: Iterable[Item]) -> list[Item]:
    """Убирает повторы по ID товара (первое вхождение побеждает)."""
    seen: dict[int, Item] = {}
    for it in items:
        seen.setdefault(it.id, it)
    return list(seen.values())


def collect(
    query: str,
    tokens: TokenStore,
    max_pages: int = 100,
    delay: tuple[float, float] = (1.0, 3.0),
    dest: str = DEFAULT_DEST,
    dump_raw: Path | None = None,
) -> list[Item]:
    """Проходит страницы выдачи, пока они не кончатся, и возвращает товары."""
    session = build_session(tokens)
    collected_at = datetime.now().isoformat(timespec="seconds")
    items: list[Item] = []
    refreshed = False
    page = 1
    while page <= max_pages:
        try:
            data = fetch_page(session, query, page, dest)
        except TokenRejected as exc:
            if refreshed:
                raise
            log.warning("Токен отклонён (%s), получаю новый через браузер", exc)
            tokens.refresh(query)
            session = build_session(tokens)
            refreshed = True
            continue
        if dump_raw:
            dump_raw.mkdir(parents=True, exist_ok=True)
            (dump_raw / f"page_{page:03d}.json").write_text(
                json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8"
            )
        page_items = parse_items(data, collected_at)
        log.info("Страница %d: найдено товаров %d", page, len(page_items))
        if not page_items:
            break
        items.extend(page_items)
        page += 1
        if page <= max_pages:
            time.sleep(random.uniform(*delay))
    unique = dedupe_items(items)
    if len(unique) != len(items):
        log.info("Дедупликация: убрано повторов %d", len(items) - len(unique))
    return unique


# ------------------------------------------------------------------ экспорт

FIELD_TITLES = {
    "id": "Артикул", "name": "Название", "brand": "Бренд", "model": "Модель",
    "memory": "Память", "color": "Цвет", "price": "Цена, ₽",
    "price_basic": "Цена без скидки, ₽", "discount_pct": "Скидка, %",
    "seller": "Продавец", "seller_id": "ID продавца", "rating": "Рейтинг",
    "feedbacks": "Отзывов", "url": "Ссылка", "collected_at": "Дата сбора",
}


def save_results(items: list[Item], path: Path) -> list[Path]:
    """Сохраняет товары в CSV и XLSX рядом друг с другом, возвращает пути."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    path.parent.mkdir(parents=True, exist_ok=True)
    names = [f.name for f in fields(Item)]
    rows = [asdict(it) for it in items]

    csv_path = path.with_suffix(".csv")
    with csv_path.open("w", newline="", encoding="utf-8-sig") as fh:  # BOM — для Excel
        writer = csv.DictWriter(fh, fieldnames=names, delimiter=";")
        writer.writeheader()
        writer.writerows(rows)

    xlsx_path = path.with_suffix(".xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "iPhone 16"
    ws.append([FIELD_TITLES[n] for n in names])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([row[n] for n in names])
    widths = {"name": 60, "url": 50, "seller": 28, "color": 18, "collected_at": 20}
    for idx, n in enumerate(names, start=1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = widths.get(n, 14)
        if n in ("price", "price_basic"):
            for cell in ws[letter][1:]:
                cell.number_format = "# ##0"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(xlsx_path)
    log.info("Сохранено: %s, %s", csv_path, xlsx_path)
    return [csv_path, xlsx_path]


def fmt_rub(value: float | None) -> str:
    """Форматирует сумму: 89 990 ₽."""
    return "—" if value is None else f"{value:,.0f} ₽".replace(",", " ")


def summarize(items: list[Item]) -> str:
    """Текстовая сводка: цены, разбивка по моделям и памяти, минимумы."""
    prices = [it.price for it in items if it.price is not None]
    lines = [f"Всего товаров: {len(items)}"]
    if prices:
        lines.append(
            f"Цена: мин {fmt_rub(min(prices))}, макс {fmt_rub(max(prices))}, "
            f"медиана {fmt_rub(statistics.median(prices))}"
        )
    by_model = Counter(it.model or "?" for it in items)
    lines.append("По моделям: " + ", ".join(
        f"{m} — {by_model[m]}" for m in sorted(by_model, key=_model_order)))
    by_mem = Counter(it.memory or "?" for it in items)
    lines.append("По памяти: " + ", ".join(
        f"{m} — {by_mem[m]}" for m in sorted(by_mem, key=_memory_order)))
    cheapest: dict[tuple[str, str], Item] = {}
    for it in items:
        key = (it.model or "?", it.memory or "?")
        if it.price is not None and (key not in cheapest or it.price < cheapest[key].price):
            cheapest[key] = it
    if cheapest:
        lines.append("Минимальная цена по модели и памяти:")
        for key in sorted(cheapest, key=lambda k: (_model_order(k[0]), _memory_order(k[1]))):
            lines.append(f"  {key[0]} {key[1]}: {fmt_rub(cheapest[key].price)}")
    return "\n".join(lines)


def _model_order(model: str) -> int:
    return MODELS.index(model) if model in MODELS else len(MODELS)


def _memory_order(memory: str) -> int:
    order = ["128 ГБ", "256 ГБ", "512 ГБ", "1 ТБ"]
    return order.index(memory) if memory in order else len(order)


# -------------------------------------------------------------------- CLI


def build_arg_parser() -> argparse.ArgumentParser:
    """Аргументы командной строки."""
    ap = argparse.ArgumentParser(description="Парсер цен на iPhone 16 с Wildberries")
    ap.add_argument("-q", "--query", default="iPhone 16", help="поисковый запрос")
    ap.add_argument("-p", "--max-pages", type=int, default=100, help="максимум страниц")
    ap.add_argument("--delay-min", type=float, default=1.0, help="мин. пауза между запросами, с")
    ap.add_argument("--delay-max", type=float, default=3.0, help="макс. пауза между запросами, с")
    ap.add_argument("-o", "--output", type=Path, default=Path("results/iphone16.xlsx"),
                    help="выходной файл (рядом создаются .xlsx и .csv)")
    ap.add_argument("--include-accessories", action="store_true",
                    help="не фильтровать аксессуары и прочие товары")
    ap.add_argument("--min-price", type=float, default=20000,
                    help="нижняя граница цены смартфона, ₽ (отсекает аксессуары)")
    ap.add_argument("--dest", default=DEFAULT_DEST, help="регион доставки WB (параметр dest)")
    ap.add_argument("--profile-dir", type=Path, default=Path(".wb_profile"),
                    help="профиль браузера для получения токена")
    ap.add_argument("--headless", action="store_true",
                    help="браузер без окна (WB часто такой не пускает)")
    ap.add_argument("--browser-channel", choices=["chrome", "msedge"],
                    help="взять установленный Chrome/Edge вместо встроенного Chromium")
    ap.add_argument("--no-browser", action="store_true",
                    help="не открывать браузер: только токен, импортированный через wb_token.py import")
    ap.add_argument("--dump-raw", type=Path, help="сохранять сырые JSON-ответы в папку")
    ap.add_argument("-v", "--verbose", action="store_true", help="подробный лог")
    return ap


def setup_logging(verbose: bool = False) -> None:
    """Настраивает формат логов."""
    logging.basicConfig(
        level=logging.DEBUG if verbose else logging.INFO,
        format="%(asctime)s %(levelname)-7s %(name)s: %(message)s",
        datefmt="%H:%M:%S",
    )


def main(argv: list[str] | None = None) -> int:
    """Точка входа CLI: сбор, фильтр, экспорт, сводка."""
    args = build_arg_parser().parse_args(argv)
    setup_logging(args.verbose)
    if args.delay_min > args.delay_max:
        args.delay_min, args.delay_max = args.delay_max, args.delay_min
    tokens = TokenStore(args.profile_dir, headless=args.headless,
                        use_browser=not args.no_browser, browser_channel=args.browser_channel)
    items = collect(args.query, tokens, args.max_pages,
                    (args.delay_min, args.delay_max), args.dest, args.dump_raw)
    items = filter_items(items, args.include_accessories, args.min_price)
    save_results(items, args.output)
    print("\n" + summarize(items))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
