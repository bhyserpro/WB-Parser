"""Тесты загрузки: повторы, обновление токена, пагинация, экспорт."""

import pytest
import requests

import wb_parser
from wb_parser import collect, fetch_page, save_results
from wb_token import TokenRejected


class Resp:
    def __init__(self, status, data=None, headers=None):
        self.status_code, self._data, self.headers = status, data, headers or {}

    def json(self):
        return self._data

    def raise_for_status(self):
        raise requests.HTTPError(self.status_code)


class FakeSession:
    def __init__(self, responses):
        self.responses, self.calls = list(responses), []
        self.headers, self.cookies = {}, requests.cookies.RequestsCookieJar()

    def get(self, url, params=None, timeout=None):
        self.calls.append(params)
        r = self.responses.pop(0)
        if isinstance(r, Exception):
            raise r
        return r


class FakeTokens:
    def __init__(self):
        self.refreshed = 0
        self.base_params = {"dest": "123", "uclusters": "2"}

    def apply(self, session, query="iPhone 16"):
        pass

    def refresh(self, query="iPhone 16"):
        self.refreshed += 1


@pytest.fixture(autouse=True)
def no_sleep(monkeypatch):
    monkeypatch.setattr(wb_parser.time, "sleep", lambda s: None)


def page(*ids):
    return {"products": [{"id": i, "name": f"iPhone 16 128 ГБ #{i}",
                          "sizes": [{"price": {"basic": 8000000, "product": 7000000}}]} for i in ids]}


def test_retry_on_429_5xx_timeout():
    s = FakeSession([Resp(429, headers={"Retry-After": "1"}), Resp(503), requests.Timeout(), Resp(200, page(1))])
    assert fetch_page(s, "iPhone 16", 1)["products"][0]["id"] == 1
    assert len(s.calls) == 4


def test_gives_up_after_retries():
    with pytest.raises(RuntimeError):
        fetch_page(FakeSession([Resp(500)] * 3), "iPhone 16", 1, retries=2)


def test_token_rejected():
    with pytest.raises(TokenRejected):
        fetch_page(FakeSession([Resp(498)]), "iPhone 16", 1)


def test_collect_paginates_dedupes_and_refreshes_token(monkeypatch, tmp_path):
    sessions = iter([
        FakeSession([Resp(200, page(1, 2)), Resp(403)]),               # токен протух на 2-й странице
        FakeSession([Resp(200, page(2, 3)), Resp(200, {"products": []})]),
    ])
    monkeypatch.setattr(wb_parser, "build_session", lambda tokens: next(sessions))
    tokens = FakeTokens()
    items = collect("iPhone 16", tokens, max_pages=10, dump_raw=tmp_path / "raw")
    assert [it.id for it in items] == [1, 2, 3] and tokens.refreshed == 1
    assert sorted(p.name for p in (tmp_path / "raw").iterdir()) == ["page_001.json", "page_002.json", "page_003.json"]

    paths = save_results(items, tmp_path / "out" / "res.xlsx")
    assert [p.suffix for p in paths] == [".csv", ".xlsx"] and all(p.stat().st_size > 0 for p in paths)
    from openpyxl import load_workbook
    ws = load_workbook(paths[1]).active
    assert ws.max_row == 4 and ws["G2"].value == 70000.0


def test_max_pages_limit(monkeypatch):
    s = FakeSession([Resp(200, page(1)), Resp(200, page(2)), Resp(200, page(3))])
    monkeypatch.setattr(wb_parser, "build_session", lambda tokens: s)
    assert len(collect("iPhone 16", FakeTokens(), max_pages=2)) == 2 and len(s.calls) == 2
    # параметры браузера сохраняются, query/page — наши
    assert s.calls[1]["uclusters"] == "2" and s.calls[1]["dest"] == "123" and s.calls[1]["page"] == "2"


def test_dest_overrides_base_params():
    s = FakeSession([Resp(200, page(1))])
    fetch_page(s, "iPhone 16", 1, dest="-1", base_params={"dest": "123", "query": "x"})
    assert s.calls[0]["dest"] == "-1" and s.calls[0]["query"] == "iPhone 16"
